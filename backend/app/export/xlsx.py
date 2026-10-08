"""Exportação em Excel: uma aba por turma (ou professor, ou sala).

`openpyxl` escreve no formato nativo, então a planilha abre com as cores, a
largura das colunas e o congelamento do cabeçalho já prontos — quem recebe não
precisa formatar nada para imprimir.
"""

from __future__ import annotations

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app import schemas as s
from app.export.grade import Quadro, faixa_texto, montar_quadros, rotulo_dia

LARGURA_HORARIO = 14
LARGURA_DIA = 24
ALTURA_AULA = 46

CINZA = "FF6B7280"
_FINA = Side(style="thin", color="FFD4D4D8")


def _titulo_de_aba(titulo: str, usados: set[str]) -> str:
    """O Excel recusa `: \\ / ? * [ ]` e nomes com mais de 31 caracteres, e ainda
    exige que sejam únicos — 'DSM 1º sem' e 'DSM 1º sem (2)' se duas turmas
    tiverem o mesmo rótulo."""
    limpo = re.sub(r"[:\\/?*\[\]]", "-", titulo).strip() or "Sem nome"
    base = limpo[:31]
    nome, n = base, 2
    while nome.casefold() in usados:
        sufixo = f" ({n})"
        nome = base[: 31 - len(sufixo)] + sufixo
        n += 1
    usados.add(nome.casefold())
    return nome


def _escrever_quadro(planilha, quadro: Quadro, cabecalho: str) -> None:
    planilha["A1"] = cabecalho
    planilha["A1"].font = Font(bold=True, size=13)
    planilha["A2"] = quadro.titulo + (f" — {quadro.subtitulo}" if quadro.subtitulo else "")
    planilha["A2"].font = Font(size=10, color=CINZA)

    linha_topo = 4
    planilha.cell(linha_topo, 1, "Horário").font = Font(bold=True)
    for coluna, dia in enumerate(quadro.dias, start=2):
        celula = planilha.cell(linha_topo, coluna, rotulo_dia(dia))
        celula.font = Font(bold=True)
        celula.alignment = Alignment(horizontal="center")

    for i, (inicio, fim) in enumerate(quadro.faixas):
        linha = linha_topo + 1 + i
        planilha.row_dimensions[linha].height = ALTURA_AULA
        rotulo = planilha.cell(linha, 1, faixa_texto(inicio, fim))
        rotulo.font = Font(size=9, color=CINZA)
        rotulo.alignment = Alignment(vertical="center")
        rotulo.border = Border(bottom=_FINA, right=_FINA)

        for coluna, dia in enumerate(quadro.dias, start=2):
            alvo = planilha.cell(linha, coluna)
            alvo.border = Border(bottom=_FINA, right=_FINA)
            alvo.alignment = Alignment(vertical="center", wrap_text=True)
            aula = quadro.aula(i, dia)
            if aula is None:
                continue
            marcas = "".join(m for m, ativo in (("🔒", aula.fixada), ("🌐", aula.remota)) if ativo)
            alvo.value = "\n".join((f"{aula.sigla}{' ' + marcas if marcas else ''}", *aula.linhas))
            # openpyxl usa ARGB; as cores vêm de `cores.py` como #RRGGBB.
            alvo.fill = PatternFill("solid", fgColor="FF" + aula.fundo[1:])
            alvo.font = Font(color="FF" + aula.texto[1:], size=9)

    planilha.column_dimensions["A"].width = LARGURA_HORARIO
    for coluna in range(2, len(quadro.dias) + 2):
        planilha.column_dimensions[get_column_letter(coluna)].width = LARGURA_DIA
    planilha.freeze_panes = planilha.cell(linha_topo + 1, 2)
    planilha.page_setup.orientation = "landscape"
    planilha.print_title_rows = f"{linha_topo}:{linha_topo}"


def gerar(visao: s.VisaoGrade, cabecalho: str) -> bytes:
    """Devolve o .xlsx pronto. Uma aba por coluna, na ordem da tela."""
    quadros = [q for q in montar_quadros(visao) if not q.vazio]
    livro = Workbook()
    livro.remove(livro.active)

    usados: set[str] = set()
    for quadro in quadros:
        aba = livro.create_sheet(_titulo_de_aba(quadro.titulo, usados))
        _escrever_quadro(aba, quadro, cabecalho)

    if not livro.sheetnames:
        # Um arquivo sem abas não abre no Excel; melhor entregar o motivo escrito.
        vazia = livro.create_sheet("Sem aulas")
        vazia["A1"] = cabecalho
        vazia["A2"] = "Esta grade ainda não tem nenhuma aula alocada."

    buffer = BytesIO()
    livro.save(buffer)
    return buffer.getvalue()
