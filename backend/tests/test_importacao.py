"""Etapa 9: importacao de cadastros por CSV e Excel."""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.importacao.campos import ValorInvalido, booleano, decimal, inteiro, opcao
from app.importacao.planilha import ErroDeArquivo, ler_tabela, normalizar_cabecalho
from app.models import Espaco, OrigemHabilitacao, Professor, TipoContrato, Unidade, VinculoUnidade

# ------------------------------------------------------------------- leitura


def _csv(texto: str, codificacao="utf-8") -> bytes:
    return texto.encode(codificacao)


def test_le_csv_com_ponto_e_virgula_e_bom_do_excel():
    conteudo = "﻿identificador;nome\nP1;Ana\n".encode()
    assert ler_tabela("x.csv", conteudo) == [{"identificador": "P1", "nome": "Ana"}]


def test_le_csv_com_virgula():
    assert ler_tabela("x.csv", _csv("identificador,nome\nP1,Ana\n")) == [
        {"identificador": "P1", "nome": "Ana"}
    ]


def test_le_csv_em_latin1_sem_estourar():
    linhas = ler_tabela("x.csv", _csv("identificador;nome\nP1;Sérgio\n", "latin-1"))
    assert linhas[0]["nome"] == "Sérgio"


def test_cabecalho_com_acento_e_espaco_vira_o_nome_da_coluna():
    assert normalizar_cabecalho("Aulas Semanais") == "aulas_semanais"
    assert normalizar_cabecalho("Pontuação") == "pontuacao"
    assert normalizar_cabecalho("tipo-contrato") == "tipo_contrato"


def test_le_xlsx_e_ignora_as_linhas_em_branco_do_fim():
    livro = Workbook()
    planilha = livro.active
    planilha.append(["identificador", "nome"])
    planilha.append(["P1", "Ana"])
    planilha.append([None, None])
    buffer = BytesIO()
    livro.save(buffer)

    assert ler_tabela("x.xlsx", buffer.getvalue()) == [{"identificador": "P1", "nome": "Ana"}]


def test_arquivo_vazio_e_formato_desconhecido_explicam_o_motivo():
    with pytest.raises(ErroDeArquivo, match="vazio"):
        ler_tabela("x.csv", b"")
    with pytest.raises(ErroDeArquivo, match="não suportado"):
        ler_tabela("x.pdf", b"%PDF")


# -------------------------------------------------------------------- campos


def test_numero_aceita_a_virgula_decimal_como_a_pontuacao_e_publicada():
    assert decimal("87,5", "pontuacao") == 87.5
    assert decimal("87.5", "pontuacao") == 87.5
    assert decimal("", "pontuacao") == 0.0


def test_booleano_aceita_sim_nao_e_vazio():
    assert booleano("sim", "x") is True
    assert booleano("NÃO", "x") is False
    assert booleano("", "x", padrao=True) is True
    with pytest.raises(ValorInvalido, match="sim ou não"):
        booleano("talvez", "x")


def test_opcao_ignora_caixa_e_acento_e_lista_os_valores_aceitos():
    assert opcao("determinado", "tipo", TipoContrato) == TipoContrato.DETERMINADO
    with pytest.raises(ValorInvalido, match="INDETERMINADO, DETERMINADO"):
        opcao("vitalicio", "tipo", TipoContrato)


def test_inteiro_recusa_texto_dizendo_o_que_veio():
    with pytest.raises(ValorInvalido, match="veio 'muitas'"):
        inteiro("muitas", "aulas_semanais")


# --------------------------------------------------------------------- rotas


def _unidade(db, nome="Fatec Taquaritinga") -> Unidade:
    return db.scalar(select(Unidade).where(Unidade.nome == nome))


def _enviar(cliente, unidade_id, tipo, texto, *, aplicar=False, nome="dados.csv"):
    return cliente.post(
        f"/importar/csv?unidade_id={unidade_id}&tipo={tipo}&aplicar={str(aplicar).lower()}",
        files={"arquivo": (nome, texto.encode(), "text/csv")},
    )


def test_lista_de_modelos_traz_as_colunas_de_cada_planilha(cliente):
    modelos = cliente.get("/importar/modelos").json()
    tipos = {m["tipo"] for m in modelos}
    assert tipos == {"professores", "disciplinas", "turmas", "espacos", "habilitacoes"}

    professores = next(m for m in modelos if m["tipo"] == "professores")
    assert any(c["nome"] == "identificador" and c["obrigatoria"] for c in professores["colunas"])


def test_modelo_em_csv_sai_com_bom_e_ponto_e_virgula(cliente):
    """E' o que faz o Excel em portugues abrir o arquivo ja' em colunas."""
    r = cliente.get("/importar/modelos/professores/arquivo?formato=csv")
    assert r.status_code == 200
    assert r.content.startswith("﻿".encode())
    assert b";" in r.content.splitlines()[0]


def test_modelo_em_xlsx_explica_cada_coluna_no_comentario(cliente):
    r = cliente.get("/importar/modelos/disciplinas/arquivo?formato=xlsx")
    planilha = load_workbook(BytesIO(r.content)).active
    assert planilha["A1"].value == "curso_sigla"
    assert planilha["A1"].comment is not None
    assert "obrigatória" in planilha["A1"].comment.text


def test_modelo_inexistente_lista_os_que_existem(cliente):
    r = cliente.get("/importar/modelos/alunos/arquivo")
    assert r.status_code == 404
    assert "professores" in r.json()["detail"]


def test_conferir_nao_grava_nada(cliente, escopo_seed, db):
    """O `aplicar=false` roda a importacao inteira e desfaz; foi aqui que o
    SAVEPOINT do SQLite escapou e a conferencia gravou de verdade."""
    unidade = _unidade(db)
    antes = db.scalar(select(func.count()).select_from(Professor))

    corpo = _enviar(
        cliente, unidade.id, "professores", "identificador;nome;pontuacao\nP900;Rita;91,5\n"
    ).json()

    assert corpo["aplicado"] is False
    assert corpo["resumo"]["criar"] == 1
    db.rollback()
    assert db.scalar(select(func.count()).select_from(Professor)) == antes


def test_aplicar_cria_o_professor_e_o_vinculo(cliente, escopo_seed, db):
    unidade = _unidade(db)
    corpo = _enviar(
        cliente,
        unidade.id,
        "professores",
        "identificador;nome;pontuacao;pont_assiduidade\nP900;Rita Nogueira;91,5;22\n",
        aplicar=True,
    ).json()

    assert corpo["aplicado"] is True
    db.rollback()
    professor = db.scalar(select(Professor).where(Professor.identificador == "P900"))
    assert professor.nome == "Rita Nogueira"
    vinculo = db.scalar(
        select(VinculoUnidade).where(
            VinculoUnidade.professor_id == professor.id,
            VinculoUnidade.unidade_id == unidade.id,
        )
    )
    assert vinculo.pontuacao == 91.5


def test_reimportar_atualiza_sem_duplicar(cliente, escopo_seed, db):
    unidade = _unidade(db)
    linha = "identificador;nome;pontuacao\nana.ribeiro;Ana Ribeiro;95,0\n"
    corpo = _enviar(cliente, unidade.id, "professores", linha, aplicar=True).json()

    assert corpo["resumo"] == {"criar": 0, "atualizar": 1, "ignorar": 0, "erro": 0}
    db.rollback()
    assert (
        db.scalar(
            select(func.count())
            .select_from(Professor)
            .where(Professor.identificador == "ana.ribeiro")
        )
        == 1
    )


def test_o_mesmo_docente_em_outra_unidade_ganha_vinculo_e_nao_um_gemeo(cliente, escopo_seed, db):
    """E' o caso do Paulo Serrano do seed: um docente, duas unidades."""
    etec = _unidade(db, "Etec Taquaritinga")
    antes = db.scalar(select(func.count()).select_from(Professor))

    linha = "identificador;nome\nana.ribeiro;Ana Ribeiro\n"
    corpo = _enviar(cliente, etec.id, "professores", linha, aplicar=True).json()

    assert corpo["resumo"]["criar"] == 1  # o vinculo e' novo
    db.rollback()
    assert db.scalar(select(func.count()).select_from(Professor)) == antes


def test_uma_linha_errada_recusa_o_arquivo_inteiro(cliente, escopo_seed, db):
    """Meio cadastro importado e' pior que nenhum."""
    unidade = _unidade(db)
    texto = "identificador;nome;tipo_contrato\nP900;Rita;INDETERMINADO\nP901;Sergio;VITALICIO\n"
    corpo = _enviar(cliente, unidade.id, "professores", texto, aplicar=True).json()

    assert corpo["aplicado"] is False
    assert corpo["resumo"]["erro"] == 1
    db.rollback()
    assert db.scalar(select(Professor).where(Professor.identificador == "P900")) is None


def test_o_erro_diz_a_linha_e_o_motivo(cliente, escopo_seed, db):
    unidade = _unidade(db)
    corpo = _enviar(
        cliente, unidade.id, "professores", "identificador;nome\nP900;Rita\n;Sem id\n"
    ).json()

    erro = next(linha for linha in corpo["linhas"] if linha["acao"] == "erro")
    assert erro["numero"] == 3  # 1 e' o cabecalho
    assert "identificador" in erro["erro"]


def test_faltando_coluna_obrigatoria_o_arquivo_nem_comeca(cliente, escopo_seed, db):
    unidade = _unidade(db)
    r = _enviar(cliente, unidade.id, "professores", "nome;pontuacao\nRita;90\n")
    assert r.status_code == 422
    assert "identificador" in r.json()["detail"]


def test_arquivo_so_com_cabecalho_avisa_em_vez_de_dizer_que_importou(cliente, escopo_seed, db):
    unidade = _unidade(db)
    r = _enviar(cliente, unidade.id, "professores", "identificador;nome\n")
    assert r.status_code == 422
    assert "cabeçalho" in r.json()["detail"]


def test_importa_espacos(cliente, escopo_seed, db):
    unidade = _unidade(db)
    texto = (
        "nome;tipo;tipo_laboratorio;capacidade\n"
        "Lab Quimica;LABORATORIO;quimica;28\n"
        "Sala 90;SALA;;45\n"
    )
    corpo = _enviar(cliente, unidade.id, "espacos", texto, aplicar=True).json()

    assert corpo["resumo"]["criar"] == 2
    db.rollback()
    lab = db.scalar(select(Espaco).where(Espaco.nome == "Lab Quimica"))
    assert lab.tipo_laboratorio == "quimica"
    assert lab.capacidade == 28


def test_importa_disciplinas_num_curso_existente(cliente, escopo_seed, db):
    unidade = _unidade(db)
    texto = (
        "curso_sigla;periodo;sigla;nome;aulas_semanais;tamanho_bloco;precisa_laboratorio\n"
        "DSM;1o sem;EST;Estatistica;4;2;nao\n"
    )
    corpo = _enviar(cliente, unidade.id, "disciplinas", texto, aplicar=True).json()
    assert corpo["resumo"]["criar"] == 1


def test_curso_que_nao_existe_vira_erro_de_linha_e_nao_500(cliente, escopo_seed, db):
    unidade = _unidade(db)
    texto = "curso_sigla;periodo;sigla;nome\nXPTO;1o sem;EST;Estatistica\n"
    corpo = _enviar(cliente, unidade.id, "disciplinas", texto).json()

    assert corpo["resumo"]["erro"] == 1
    assert "XPTO" in corpo["linhas"][0]["erro"]


def test_importa_habilitacoes(cliente, escopo_seed, db):
    unidade = _unidade(db)
    texto = (
        "professor_identificador;curso_sigla;periodo;disciplina_sigla;origem\n"
        "bruno.carvalho;DSM;1o sem;ALP;CONCURSO\n"
    )
    corpo = _enviar(cliente, unidade.id, "habilitacoes", texto, aplicar=True).json()

    assert corpo["resumo"]["erro"] == 0
    db.rollback()
    from app.models import Habilitacao

    habilitacao = db.scalar(
        select(Habilitacao)
        .join(Professor)
        .where(Professor.identificador == "bruno.carvalho")
        .order_by(Habilitacao.id.desc())
    )
    assert habilitacao.origem == OrigemHabilitacao.CONCURSO


def test_importa_xlsx_tambem(cliente, escopo_seed, db):
    unidade = _unidade(db)
    livro = Workbook()
    planilha = livro.active
    planilha.append(["nome", "tipo", "capacidade"])
    planilha.append(["Sala 91", "SALA", 30])
    buffer = BytesIO()
    livro.save(buffer)

    r = cliente.post(
        f"/importar/csv?unidade_id={unidade.id}&tipo=espacos&aplicar=true",
        files={"arquivo": ("salas.xlsx", buffer.getvalue(), "application/octet-stream")},
    )
    assert r.json()["resumo"]["criar"] == 1
